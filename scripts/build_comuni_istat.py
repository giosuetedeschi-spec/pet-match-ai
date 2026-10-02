"""Build the current municipality CSV from official ISTAT 2026 sources."""

import argparse
import csv
import io
import zipfile
from pathlib import Path

import shapefile
from openpyxl import load_workbook
from pyproj import Transformer
from shapely.geometry import shape
from shapely.ops import unary_union


BOUNDARY_MERGES = {
    # The January 1 boundary release predates these 2026 administrative changes.
    '018094': ('018094', '018082'),  # Lirio incorporated into Montalto Pavese.
    '024129': ('024027', '024071'),  # Castegnero and Nanto became Castegnero Nanto.
}
OUTPUT_FIELDS = (
    'istat_code', 'name', 'province_code', 'province', 'province_abbreviation',
    'region', 'latitude', 'longitude',
)


def build(boundaries_zip: Path, municipalities_xlsx: Path, output_csv: Path) -> int:
    with zipfile.ZipFile(boundaries_zip) as archive:
        shape_path = next(
            name for name in archive.namelist()
            if name.endswith('/Com01012026_g_WGS84.shp')
        )
        base = shape_path[:-4]
        reader = shapefile.Reader(
            shp=io.BytesIO(archive.read(shape_path)),
            shx=io.BytesIO(archive.read(base + '.shx')),
            dbf=io.BytesIO(archive.read(base + '.dbf')),
            encoding='utf-8',
        )
        geometries = {
            record['PRO_COM_T']: shape_record.shape
            for shape_record in reader.iterShapeRecords()
            for record in (shape_record.record,)
        }
    geometries = {code: shape(geometry.__geo_interface__) for code, geometry in geometries.items()}

    workbook = load_workbook(municipalities_xlsx, read_only=True, data_only=True)
    sheet = workbook['CODICI al 21_02_2026']
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)

    def column_index(prefix: str) -> int:
        for index, header in enumerate(headers):
            if str(header).replace('\n', ' ').startswith(prefix):
                return index
        raise ValueError(f'ISTAT spreadsheet is missing a column starting with {prefix!r}')

    columns = {
        'code': column_index('Codice Comune formato alfanumerico'),
        'name': column_index('Denominazione in italiano'),
        'province_code': column_index("Codice dell'Unità territoriale sovracomunale"),
        'province': column_index("Denominazione dell'Unità territoriale sovracomunale"),
        'province_abbreviation': column_index('Sigla automobilistica'),
        'region': column_index('Denominazione Regione'),
    }
    transformer = Transformer.from_crs('EPSG:32632', 'EPSG:4326', always_xy=True)

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with output_csv.open('w', newline='', encoding='utf-8-sig') as output_file:
        writer = csv.DictWriter(output_file, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()

        for row in rows:
            code = row[columns['code']]
            source_codes = BOUNDARY_MERGES.get(code, (code,))
            missing = [source_code for source_code in source_codes if source_code not in geometries]
            if missing:
                raise ValueError(f'Missing boundary geometry for {code}: {", ".join(missing)}')

            municipality_geometry = unary_union([geometries[source_code] for source_code in source_codes])
            centroid = municipality_geometry.centroid
            longitude, latitude = transformer.transform(centroid.x, centroid.y)
            writer.writerow({
                'istat_code': code,
                'name': row[columns['name']],
                'province_code': str(row[columns['province_code']]).zfill(3),
                'province': row[columns['province']],
                'province_abbreviation': row[columns['province_abbreviation']],
                'region': row[columns['region']],
                'latitude': f'{latitude:.6f}',
                'longitude': f'{longitude:.6f}',
            })
            count += 1

    workbook.close()
    return count


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--boundaries', required=True, type=Path, help='ISTAT Limiti01012026_g.zip')
    parser.add_argument('--municipalities', required=True, type=Path, help='ISTAT Elenco-comuni-italiani.xlsx')
    parser.add_argument('--output', default=Path('data/comuni_istat_2026.csv'), type=Path)
    args = parser.parse_args()
    count = build(args.boundaries, args.municipalities, args.output)
    print(f'Wrote {count} municipalities to {args.output}')


if __name__ == '__main__':
    main()
